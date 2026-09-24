"""Read a supplier workbook; create a private JSON import manifest, never edit the workbook."""
import argparse, json, re
from pathlib import Path
import openpyxl

p=argparse.ArgumentParser()
p.add_argument('workbook');p.add_argument('--output',default='migration-data/supplier-import.json')
args=p.parse_args()
book=openpyxl.load_workbook(args.workbook,data_only=True)
groups={}; counts={}; excluded=[]
def value(v):
    if v is None: return ''
    if isinstance(v,float) and v.is_integer(): return str(int(v))
    return str(v).strip()
for sheet in book:
    category={'Computer PeripheralsAccessories':'Computer Peripherals/Accessories','PrintingFabricating':'Printing/Fabricating'}.get(sheet.title.strip(),sheet.title.strip())
    if re.search(r'fire\s*(?:and|&)\s*safety|document\s+management',category,re.I):
        excluded.append(category);continue
    headers=[value(sheet.cell(1,c).value) for c in range(1,8)]
    if headers!=['#','Name','Address','Item','Email Address','Contact Number','Contact Person']:
        raise ValueError(f'Unexpected headers in {category}: {headers}')
    counts[category]=0
    for number,row in enumerate(sheet.iter_rows(min_row=2,max_col=7,values_only=True),2):
        _,name,address,item,email,phone,contact=map(value,row)
        if not name:
            if any([address,item,email,phone,contact]): raise ValueError(f'Missing supplier name: {category}:{number}')
            continue
        key=' '.join(name.split()).casefold()
        entry=groups.setdefault(key,{'name':name,'rows':[]})
        entry['rows'].append({'sheet':category,'row':number,'address':address,'item':item,'email':email,'phone':phone,'contact':contact})
        counts[category]+=1
records=[]
for entry in groups.values():
    rows=entry['rows']
    def joined(key): return '\n'.join(dict.fromkeys(r[key] for r in rows if r[key]))
    records.append({'name':entry['name'],'category':joined('sheet').replace('\n',', '),'contact':joined('contact'),'email':joined('email'),'phone':joined('phone'),'address':joined('address'),'notes':joined('item'),'source':'Supplier tracker 2026-09-23; '+', '.join(f"{r['sheet']} row {r['row']}" for r in rows),'rows':rows})
output={'source':'https://docs.google.com/spreadsheets/d/1cSNeYwaYL9jyH4yR9OVEDtF-qhAy7JNEbHaOzIuSrXM/edit','excluded':excluded,'counts':counts,'suppliers':records}
Path(args.output).parent.mkdir(parents=True,exist_ok=True)
Path(args.output).write_text(json.dumps(output,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'source_rows':sum(counts.values()),'unique_suppliers':len(records),'excluded':excluded,'counts':counts,'max_field_lengths':{k:max((len(r[k]) for r in records),default=0) for k in ['category','contact','email','phone','address','notes']}}))
