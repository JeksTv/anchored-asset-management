#!/usr/bin/env python3
"""Preview or apply the 20-tab Google Sheets hardware migration.

The Google Sheet must first be exported as XLSX. Preview is the default. Use
--apply only after reviewing the CSV report. Re-running is idempotent by tag.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import sqlite3
import unicodedata
from datetime import date, datetime
from pathlib import Path

import openpyxl

TABS = [
    "MONITOR", "MOUSE", "DESKTOP", "TABLET", "HEADSET", "KEYBOARD",
    "WEBCAM", "HDMI", "VGA", "EXTERNAL HDD", "SSD", "RAM",
    "WIFI ADAPTER", "PRINTER", "IPAD", "SMART TV", "CONFERENCE MIC",
    "TAPO CAMERA", "VGA-HDMI", "NOTEBOOK COOLER",
]

CATEGORY = {
    "MONITOR": "Monitor", "MOUSE": "Mouse", "DESKTOP": "Desktop",
    "TABLET": "Tablet", "HEADSET": "Headset", "KEYBOARD": "Keyboard",
    "WEBCAM": "Webcam", "HDMI": "HDMI cable", "VGA": "VGA cable",
    "EXTERNAL HDD": "External HDD", "SSD": "SSD", "RAM": "RAM",
    "WIFI ADAPTER": "WiFi adapter", "PRINTER": "Printer", "IPAD": "iPad",
    "SMART TV": "Smart TV", "CONFERENCE MIC": "Conference microphone",
    "TAPO CAMERA": "Tapo camera", "VGA-HDMI": "VGA to HDMI adapter",
    "NOTEBOOK COOLER": "Notebook cooler",
}

GENERIC_CUSTODIANS = {
    "available", "it", "it room", "ro", "warehouse", "server room",
    "server room at ro", "server room at lp", "bgc hub", "tektite hub",
    "alabang hub", "vcis", "el/lp", "room", "records 1",
}


def clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def norm(value: str) -> str:
    value = unicodedata.normalize("NFKD", clean(value)).encode("ascii", "ignore").decode()
    return " ".join(re.findall(r"[a-z0-9]+", value.lower()))


def iso_date(value) -> str:
    if value is None or clean(value) == "":
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = clean(value).replace("Sept.", "Sep").replace("Sept ", "Sep ")
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%b %d, %Y", "%B %d, %Y", "%b. %d, %Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return text


def stable_id(prefix: str, key: str) -> str:
    return "migration-" + hashlib.sha256(f"asset-tabs:{prefix}:{key.lower()}".encode()).hexdigest()


def row_value(row, indexes: dict[str, list[int]], header: str, occurrence: int = 0):
    found = indexes.get(header.lower(), [])
    if occurrence >= len(found):
        return ""
    index = found[occurrence]
    return row[index] if index < len(row) else ""


def employee_match(custodian: str, employees: list[sqlite3.Row], aliases: dict[str, str]):
    target = norm(custodian)
    if not target or target in GENERIC_CUSTODIANS or any(x in target for x in (" hub", "room", "warehouse", "server")):
        return None, "location"
    alias_code = aliases.get(target)
    if alias_code:
        candidates = [e for e in employees if e["code"].lower() == alias_code.lower()]
        if len(candidates) == 1:
            return candidates[0], "reviewed identity"
    exact = [e for e in employees if norm(e["name"]) == target]
    if len(exact) == 1:
        return exact[0], "exact"
    tokens = target.split()
    if len(tokens) >= 2:
        first, last = tokens[0], tokens[-1]
        candidates = [e for e in employees if first in norm(e["name"]).split() and last in norm(e["name"]).split()]
        if len(candidates) == 1:
            return candidates[0], "unique first/last"
    return None, "review"


def state_and_condition(source_condition: str, age: str):
    value = norm(source_condition)
    if any(x in value for x in ("not working", "damaged", "defective", "broken")):
        return "Retired", "Damaged"
    if "brandnew" in norm(age) or norm(age) == "new":
        return "Ready", "New"
    return "Ready", "Good"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--database", type=Path, default=Path("data/anchored.sqlite"))
    parser.add_argument("--report", type=Path, default=Path("migration-data/asset-tab-migration-report.csv"))
    parser.add_argument("--decisions", type=Path, default=Path("migration-data/laptop-decisions.json"))
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    if not args.workbook.is_file():
        raise SystemExit(f"Workbook not found: {args.workbook}")
    if not args.database.is_file():
        raise SystemExit(f"Database not found: {args.database}")

    workbook = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True)
    missing = [name for name in TABS if name not in workbook.sheetnames]
    if missing:
        raise SystemExit("Missing tabs: " + ", ".join(missing))

    db = sqlite3.connect(args.database, timeout=30)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA foreign_keys=ON")
    employees = db.execute("SELECT id,code,name,email,status FROM employees").fetchall()
    aliases: dict[str, str] = {}
    if args.decisions.is_file():
        decisions = json.loads(args.decisions.read_text(encoding="utf-8-sig")).get("decisions", [])
        grouped: dict[str, set[str]] = {}
        for item in decisions:
            if item.get("decision") == "confirmed" and item.get("employee_id") and item.get("source_recipient"):
                grouped.setdefault(norm(item["source_recipient"]), set()).add(clean(item["employee_id"]))
        aliases = {name: next(iter(codes)) for name, codes in grouped.items() if len(codes) == 1}
    existing = {norm(r["tag"]): r for r in db.execute("SELECT id,tag,serial,kind,state,details FROM assets")}
    seen_tags: set[str] = set()
    staged = []

    for tab in TABS:
        rows = workbook[tab].iter_rows(values_only=True)
        headers = [clean(v).lower() for v in next(rows)]
        indexes: dict[str, list[int]] = {}
        for index, header in enumerate(headers):
            if header:
                indexes.setdefault(header, []).append(index)
        for source_row, row in enumerate(rows, start=2):
            tag = clean(row_value(row, indexes, "Control Number"))
            if not tag:
                continue
            key = norm(tag)
            if key in seen_tags:
                raise SystemExit(f"Duplicate asset tag in workbook: {tag}")
            seen_tags.add(key)
            source_asset = clean(row_value(row, indexes, "Asset")) or CATEGORY[tab]
            model = clean(row_value(row, indexes, "Model"))
            serial = clean(row_value(row, indexes, "Serial Number"))
            if norm(serial) in {"na", "n a", "none", "-"}:
                serial = ""
            age = clean(row_value(row, indexes, "Brandnew / Old"))
            condition_raw = clean(row_value(row, indexes, "Condition", 0))
            released = iso_date(row_value(row, indexes, "Date Released"))
            custodian = clean(row_value(row, indexes, "Released to"))
            employee, confidence = employee_match(custodian, employees, aliases)
            active_custody = bool(custodian)
            state, condition = state_and_condition(condition_raw, age)
            details = {
                "category": CATEGORY[tab], "model": model, "condition": condition,
                "source_sheet": tab, "source_row": source_row,
                "source_asset_name": source_asset, "source_new_or_old": age,
                "source_condition": condition_raw, "released_to": custodian,
                "date_released": released,
            }
            details = {k: v for k, v in details.items() if v not in ("", None)}
            if active_custody and not employee:
                details["custodian"] = custodian
                details["custody_type"] = "location" if confidence == "location" else "unmatched source custodian"
            old = existing.get(key)
            conflict = ""
            action = "create"
            if old:
                if old["kind"] != "Hardware":
                    conflict = "existing tag is not hardware"
                elif clean(old["serial"]).lower() != serial.lower():
                    conflict = "existing tag has different serial"
                else:
                    action = "match existing"
            if employee and employee["status"] not in ("Active", "Onboarding"):
                employee = None
                confidence = "departed employee; review"
            staged.append({
                "tab": tab, "source_row": source_row, "tag": tag, "name": source_asset,
                "serial": serial, "state": state, "details": details, "released_to": custodian,
                "employee_id": employee["id"] if employee and active_custody else "",
                "employee_code": employee["code"] if employee and active_custody else "",
                "employee_name": employee["name"] if employee and active_custody else "",
                "match": confidence if active_custody else "not currently deployed",
                "assigned_at": released, "active_custody": active_custody,
                "action": action, "conflict": conflict,
            })

    args.report.parent.mkdir(parents=True, exist_ok=True)
    fields = ["tab", "source_row", "tag", "name", "serial", "state", "released_to", "employee_code", "employee_name", "match", "assigned_at", "active_custody", "action", "conflict"]
    with args.report.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows({k: row[k] for k in fields} for row in staged)

    conflicts = [row for row in staged if row["conflict"]]
    if conflicts:
        raise SystemExit(f"{len(conflicts)} conflicts found. Review {args.report}; no changes applied.")

    if args.apply:
        new_asset_count = sum(1 for row in staged if row["action"] == "create")
        updated_asset_count = sum(
            1 for row in staged
            if row["action"] != "create"
            and json.loads(existing[norm(row["tag"])]["details"] or "{}").get("source_sheet") == row["tab"]
            and existing[norm(row["tag"])]["details"] != json.dumps(row["details"], ensure_ascii=False)
        )
        new_assignment_count = 0
        for row in staged:
            if not row["employee_id"]:
                continue
            old = existing.get(norm(row["tag"]))
            asset_id = old["id"] if old else stable_id("asset", row["tag"])
            if not db.execute("SELECT 1 FROM assignments WHERE asset_id=? AND resolved_at IS NULL", (asset_id,)).fetchone():
                new_assignment_count += 1
        if new_asset_count == 0 and updated_asset_count == 0 and new_assignment_count == 0:
            print(f"No changes required; all {len(staged)} source assets are already present.")
            return
        backup_dir = Path("backups") / datetime.now().strftime("%Y-%m-%dT%H-%M-%S-asset-tabs")
        backup_dir.mkdir(parents=True, exist_ok=False)
        backup_db = sqlite3.connect(backup_dir / "anchored.sqlite")
        db.backup(backup_db)
        backup_db.close()
        readme = "Pre-import database backup for the 20-tab hardware migration. Contains personal data and login hashes; restrict access.\n"
        (backup_dir / "README.txt").write_text(readme, encoding="utf-8")
        uploads = args.database.parent / "uploads"
        if uploads.exists():
            shutil.copytree(uploads, backup_dir / "uploads")

        now = datetime.now().astimezone().isoformat()
        db.execute("BEGIN IMMEDIATE")
        try:
            for row in staged:
                old = existing.get(norm(row["tag"]))
                asset_id = old["id"] if old else stable_id("asset", row["tag"])
                if not old:
                    db.execute(
                        "INSERT INTO assets(id,tag,name,kind,serial,seats,state,details) VALUES(?,?,?,'Hardware',?,1,?,?)",
                        (asset_id, row["tag"], row["name"], row["serial"], row["state"], json.dumps(row["details"], ensure_ascii=False)),
                    )
                elif json.loads(old["details"] or "{}").get("source_sheet") == row["tab"]:
                    db.execute(
                        "UPDATE assets SET name=?,serial=?,state=?,details=? WHERE id=?",
                        (row["name"], row["serial"], row["state"], json.dumps(row["details"], ensure_ascii=False), asset_id),
                    )
                event_id = stable_id("asset-event", row["tag"])
                db.execute(
                    "INSERT OR IGNORE INTO asset_events(id,asset_id,message,snapshot,created_at) VALUES(?,?,?,?,?)",
                    (event_id, asset_id, f"Imported from {row['tab']} tab of ITD - 2.0 ASSET MONITORING TRACKER.", json.dumps(row["details"], ensure_ascii=False), now),
                )
                if row["employee_id"]:
                    active = db.execute("SELECT employee_id FROM assignments WHERE asset_id=? AND resolved_at IS NULL", (asset_id,)).fetchone()
                    if active and active["employee_id"] != row["employee_id"]:
                        raise RuntimeError(f"Custody conflict for {row['tag']}")
                    if not active:
                        assigned_at = row["assigned_at"] or now[:10]
                        db.execute(
                            "INSERT INTO assignments(id,employee_id,asset_id,identifier,assigned_at) VALUES(?,?,?,?,?)",
                            (stable_id("assignment", row["tag"]), row["employee_id"], asset_id, f"Migrated from {row['tab']} tracker; existing custody, not a new handover.", assigned_at + "T00:00:00.000Z"),
                        )
                        db.execute(
                            "INSERT OR IGNORE INTO events(id,employee_id,message,created_at) VALUES(?,?,?,?)",
                            (stable_id("employee-event", row["tag"]), row["employee_id"], f"{row['tag']} imported from {row['tab']} with {row['match']} employee match.", now),
                        )
            violations = db.execute("PRAGMA foreign_key_check").fetchall()
            if violations:
                raise RuntimeError("Foreign-key validation failed")
            db.commit()
        except Exception:
            db.rollback()
            raise
        print(f"Applied {new_asset_count} new assets; updated {updated_asset_count} matching imported assets; created {new_assignment_count} employee assignments.")
        print(f"Backup: {backup_dir}")
    else:
        print(f"Previewed {len(staged)} assets across {len(TABS)} tabs.")
        print(f"New: {sum(1 for r in staged if r['action']=='create')}; existing: {sum(1 for r in staged if r['action']!='create')}; employee assignments: {sum(1 for r in staged if r['employee_id'])}; deployed to location/unmatched custodian: {sum(1 for r in staged if r['active_custody'] and not r['employee_id'])}.")
        print(f"Report: {args.report}")


if __name__ == "__main__":
    main()
